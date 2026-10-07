"""Sinh FO_flows.md + FO_tc_mapping.xlsx từ file test case gốc của khách hàng.

Chạy lại mỗi khi khách gửi phiên bản mới:
    python build_fo_mapping.py "<đường dẫn file xlsx gốc>"

File gốc KHÔNG bị sửa. Cột trạng thái automation trong mapping cũ được giữ lại
(khớp theo TC No) nếu FO_tc_mapping.xlsx đã tồn tại.
"""
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path(__file__).resolve().parent
MAPPING = OUT_DIR / "FO_tc_mapping.xlsx"
FLOWS_MD = OUT_DIR / "FO_flows.md"

# --- Luồng happy (E2E) -------------------------------------------------------
# Biến thể của mỗi luồng = số cấp duyệt x có/không ký số pháp nhân
FLOWS = [
    ("HF-01", "Tạo mới → Gửi duyệt → Duyệt → Ký số → Đã gửi ngân hàng",
     {1: "1 cấp, không ký số PN", 2: "1 cấp, có ký số PN", 3: "2 cấp, không ký số PN", 4: "2 cấp, có ký số PN"}),
    ("HF-02", "Tạo mới → Lưu nháp → Mở lại, Gửi duyệt → Duyệt → Ký số → Đã gửi ngân hàng",
     {7: "1 cấp, không ký số PN", 5: "1 cấp, có ký số PN", 8: "2 cấp, không ký số PN", 6: "2 cấp, có ký số PN"}),
    ("HF-03", "Tạo mới → Gửi duyệt → Trả lại (Yêu cầu chỉnh sửa) → Maker sửa, gửi lại → Duyệt → Ký số → Đã gửi ngân hàng",
     {14: "1 cấp, không ký số PN", 15: "1 cấp, có ký số PN", 16: "2 cấp, không ký số PN", 17: "2 cấp, có ký số PN"}),
]
NEGATIVE_E2E = {
    9: "Hủy giao dịch từ bản nháp → Đã hủy",
    13: "Trả lại → Maker hủy → Đã hủy",
    10: "Từ chối tại Kế toán trưởng → Đã hủy",
    11: "Từ chối tại Chủ tài khoản → Từ chối",
    12: "Bước Ký số pháp nhân KHÔNG có nút Từ chối",
}

# --- Các bước trong luồng (step) --------------------------------------------
# (step_id, tên bước, flows áp dụng, rule: màn hình gốc hoặc dải TC No trong "Thêm mới")
STEPS = [
    ("S01", "Danh sách & tìm kiếm giao dịch", "HF-01/02/03", {"screen": "Danh sách và tìm kiếm giao dịch"}),
    ("S02", "Thêm mới – quyền truy cập, điều hướng wizard, tích hợp CIF/GLIMS/Core, hạn mức ngày, Gửi duyệt/Lưu nháp", "HF-01/02/03", {"range": (154, 180)}),
    ("S03", "Thêm mới – rule chung các trường (khoảng trắng, bắt buộc, độ dài, mặc định)", "HF-01/02/03", {"range": (181, 232)}),
    ("S04", "Thêm mới – Thông tin khách hàng", "HF-01/02/03", {"range": (233, 246)}),
    ("S05", "Thêm mới – Thông tin LC: loại LC, BIC, ngày, ngân hàng, bên phát hành, người hưởng", "HF-01/02/03", {"range": (247, 321)}),
    ("S06", "Thêm mới – Giá trị LC, dung sai, nơi có giá trị", "HF-01/02/03", {"range": (322, 331)}),
    ("S07", "Thêm mới – Phương thức thanh toán, hối phiếu, kỳ hạn", "HF-01/02/03", {"range": (332, 351)}),
    ("S08", "Thêm mới – Giao hàng & hàng hoá", "HF-01/02/03", {"range": (352, 387)}),
    ("S09", "Thêm mới – Điều khoản bổ sung", "HF-01/02/03", {"range": (388, 403)}),
    ("S10", "Thêm mới – Phí & thời hạn xuất trình chứng từ", "HF-01/02/03", {"range": (404, 441)}),
    ("S11", "Thêm mới – Thông tin khác, nguồn tiền, lãi suất, phương thức trả nợ", "HF-01/02/03", {"range": (442, 502)}),
    ("S12", "Thêm mới – Chứng từ yêu cầu (Documents required)", "HF-01/02/03", {"range": (689, 707)}),
    ("S13", "Thêm mới – Additional conditions, LC xác nhận/chuyển nhượng, phí discounting", "HF-01/02/03", {"range": (708, 745)}),
    ("S14", "Thêm mới – Cam kết bổ sung chứng từ (tạm nhập tái xuất / chuyển khẩu / trả trước)", "HF-01/02/03", {"range": (746, 775)}),
    ("S15", "Thêm mới – Điều khoản & điều kiện, Hồ sơ đính kèm (upload)", "HF-01/02/03", {"range": (503, 532)}),
    ("S16", "Thêm mới – Nút Lưu nháp / Tiếp tục / Quay lại / Lưu template / Gửi duyệt", "HF-01/02/03", {"range": (533, 540)}),
    ("S17", "Chỉnh sửa (bản nháp / Yêu cầu chỉnh sửa)", "HF-02, HF-03", {"screen": "Chỉnh sửa"}),
    ("S18", "Xem chi tiết", "HF-01/02/03", {"screen": "Xem chi tiết"}),
    ("S19", "Phê duyệt (Kế toán phó / Kế toán trưởng, 2FA)", "HF-01/02/03", {"screen": "Phê duyệt"}),
    ("S20", "Ký số cá nhân (Chủ tài khoản)", "HF-01/02/03", {"screen": "Ký số cá nhân"}),
    ("S21", "Ký số pháp nhân", "HF-01/02/03 (biến thể có ký số PN)", {"screen": "Ký số pháp nhân"}),
    ("N01", "Từ chối", "Luồng lỗi", {"screen": "Từ chối"}),
    ("N02", "Hủy đề nghị", "Luồng lỗi", {"screen": "Hủy đề nghị"}),
    ("N03", "Xóa đề nghị", "Luồng lỗi", {"screen": "Xóa đề nghị"}),
    ("U01", "Sao chép đề nghị", "Tiện ích", {"screen": "Sao chép"}),
    ("U02", "Quản lý template", "Tiện ích", {"screen": "Quản lý template"}),
    ("P01", "Bổ sung hồ sơ sau phát hành", "Sau phát hành (cần BO phát hành xong)", {"screen": "Bổ sung hồ sơ sau phát hành"}),
    ("P02", "Dialog – Nhập thông tin hồ sơ hậu kiểm", "Sau phát hành (cần BO phát hành xong)", {"screen": "Dialog - Nhập thông tin hồ sơ hậu kiểm"}),
]

# Gợi ý khả năng automation – chỉ là đề xuất, cần xác nhận với team
MANUAL_HINTS = [
    ("ký số", "Cần xác nhận: ký số (USB token/HSM?)"),
    ("2fa", "Cần xác nhận: 2FA/OTP"),
    ("otp", "Cần xác nhận: 2FA/OTP"),
    ("glims", "Cần xác nhận: tích hợp GLIMS"),
    ("core", "Cần xác nhận: tích hợp Core"),
    ("api", "Cần xác nhận: gọi API ngoài"),
    ("email", "Cần xác nhận: kiểm tra email"),
]

HEADERS = ["TC No", "Type", "Flow / Step ID", "Tên bước", "Áp dụng cho luồng", "Loại",
           "Màn hình (gốc)", "Function", "Điều kiện 1", "Điều kiện 2", "Summary",
           "Data", "Action", "Expected Result", "Priority",
           "Gợi ý automation", "Automatable?", "Katalon Test Case", "Trạng thái", "Ghi chú", "Note (gốc)"]
TRACK_COLS = ["Automatable?", "Katalon Test Case", "Trạng thái", "Ghi chú"]


def tc_no(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def load_fo_rows(src):
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    out = []
    for r in rows[1:]:
        r = [("" if c is None else str(c).strip()) for c in r[:18]]
        if not any(r[:12]):
            continue
        typ, screen = r[1], r[4]
        # 4 dòng E2E bị trống cột Type (TC 5, 7, 8, 12) – vẫn là FO
        if typ.startswith("FO") or (not typ and screen == "Tạo đề nghị phát hành"):
            out.append(r)
    return out


def classify(r):
    no, screen = tc_no(r[0]), r[4]
    for fid, _, variants in FLOWS:
        if no in variants:
            return fid, variants[no], "Happy – E2E", "Kiểm tra luồng"
    if no in NEGATIVE_E2E:
        return "NE", NEGATIVE_E2E[no], "Negative – E2E", "Kiểm tra luồng"
    for sid, name, applies, rule in STEPS:
        if "screen" in rule and screen == rule["screen"]:
            return sid, name, "Edge case", applies
        if "range" in rule and screen == "Thêm mới" and rule["range"][0] <= no <= rule["range"][1]:
            return sid, name, "Edge case", applies
    return "UNMAPPED", "", "Chưa phân loại", ""


def hint(r):
    text = " ".join(r[2:12]).lower()
    hits = sorted({h for k, h in MANUAL_HINTS if k in text})
    return "; ".join(hits)


def load_existing_tracking():
    if not MAPPING.exists():
        return {}
    ws = openpyxl.load_workbook(MAPPING)["Mapping"]
    hdr = [c.value for c in ws[1]]
    idx = {h: hdr.index(h) for h in TRACK_COLS + ["TC No"] if h in hdr}
    keep = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        keep[str(row[idx["TC No"]])] = {h: row[idx[h]] for h in TRACK_COLS if h in idx}
    return keep


def write_xlsx(records):
    old = load_existing_tracking()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapping"
    ws.append(HEADERS)
    for rec in records:
        prev = old.get(str(rec["TC No"]), {})
        for h in TRACK_COLS:
            rec[h] = prev.get(h) or rec.get(h, "")
        ws.append([rec.get(h, "") for h in HEADERS])

    head_fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    fills = {"Happy – E2E": "C6EFCE", "Negative – E2E": "FFC7CE", "Chưa phân loại": "FFEB9C"}
    type_col = HEADERS.index("Loại") + 1
    for row in ws.iter_rows(min_row=2):
        color = fills.get(row[type_col - 1].value)
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if color:
                c.fill = PatternFill("solid", fgColor=color)
    widths = {"TC No": 7, "Type": 8, "Flow / Step ID": 9, "Tên bước": 32, "Áp dụng cho luồng": 16, "Loại": 13,
              "Màn hình (gốc)": 18, "Function": 28, "Điều kiện 1": 36, "Điều kiện 2": 18, "Summary": 36,
              "Data": 20, "Action": 50, "Expected Result": 50, "Priority": 9, "Gợi ý automation": 26,
              "Automatable?": 13, "Katalon Test Case": 30, "Trạng thái": 13, "Ghi chú": 25, "Note (gốc)": 12}
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 15)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    n = ws.max_row
    for col, opts in (("Automatable?", '"Yes,No,Partial,TBD"'),
                      ("Trạng thái", '"Not started,In progress,Done,Blocked,Manual only"')):
        dv = DataValidation(type="list", formula1=opts, allow_blank=True)
        L = get_column_letter(HEADERS.index(col) + 1)
        dv.add(f"{L}2:{L}{n}")
        ws.add_data_validation(dv)
    for rec_row in range(2, n + 1):
        cell = ws.cell(rec_row, HEADERS.index("Trạng thái") + 1)
        if not cell.value:
            cell.value = "Not started"

    # Sheet tổng hợp – dùng công thức để tự cập nhật khi đổi trạng thái
    s = wb.create_sheet("Tổng hợp")
    s.append(["Flow / Step ID", "Tên bước", "Áp dụng cho luồng", "Tổng TC", "Critical/Highest", "High", "Done", "% Done"])
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    idc = get_column_letter(HEADERS.index("Flow / Step ID") + 1)
    prc = get_column_letter(HEADERS.index("Priority") + 1)
    stc = get_column_letter(HEADERS.index("Trạng thái") + 1)
    order = [f[0] for f in FLOWS] + ["NE"] + [st[0] for st in STEPS] + ["UNMAPPED"]
    groups = {}
    for rec in records:
        groups.setdefault(rec["Flow / Step ID"], (rec["Tên bước"] if rec["Loại"] == "Edge case" else rec["Loại"], rec["Áp dụng cho luồng"]))
    for i, gid in enumerate([g for g in order if g in groups], start=2):
        name, applies = groups[gid]
        rng = f"Mapping!${idc}$2:${idc}${n}"
        s.append([gid, name, applies,
                  f'=COUNTIF({rng},A{i})',
                  f'=COUNTIFS({rng},A{i},Mapping!${prc}$2:${prc}${n},"Critical")+COUNTIFS({rng},A{i},Mapping!${prc}$2:${prc}${n},"Highest")',
                  f'=COUNTIFS({rng},A{i},Mapping!${prc}$2:${prc}${n},"High")',
                  f'=COUNTIFS({rng},A{i},Mapping!${stc}$2:${stc}${n},"Done")',
                  f'=IF(D{i}=0,0,G{i}/D{i})'])
        s.cell(i, 8).number_format = "0%"
    last = s.max_row
    s.append(["TOTAL", "", "", f"=SUM(D2:D{last})", f"=SUM(E2:E{last})", f"=SUM(F2:F{last})",
              f"=SUM(G2:G{last})", f"=IF(D{last+1}=0,0,G{last+1}/D{last+1})"])
    s.cell(last + 1, 8).number_format = "0%"
    for c in s[last + 1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (14, 70, 30, 10, 16, 8, 8, 8)):
        s.column_dimensions[col].width = w
    s.freeze_panes = "A2"
    wb.move_sheet("Tổng hợp", offset=-1)
    wb.save(MAPPING)


def ids_compact(nums):
    nums = sorted(nums)
    out, start, prev = [], None, None
    for x in nums:
        if start is None:
            start = prev = x
        elif x == prev + 1:
            prev = x
        else:
            out.append(f"{start}" if start == prev else f"{start}–{prev}")
            start = prev = x
    if start is not None:
        out.append(f"{start}" if start == prev else f"{start}–{prev}")
    return ", ".join(out)


def write_md(records, src_name):
    by = OrderedDict()
    for rec in records:
        by.setdefault(rec["Flow / Step ID"], []).append(rec)

    def prio(recs):
        hi = sum(1 for r in recs if r["Priority"] in ("Critical", "Highest"))
        h = sum(1 for r in recs if r["Priority"] == "High")
        return hi, h

    L = []
    L.append("# FO – Phát hành LC: bản đồ luồng & edge case\n")
    L.append(f"> Sinh tự động bởi `build_fo_mapping.py` từ `{src_name}`. **Không sửa tay** – sửa script rồi chạy lại.")
    L.append("> Chi tiết từng TC (Action/Expected) và theo dõi tiến độ automation: `FO_tc_mapping.xlsx`.\n")
    total = len(records)
    e2e = sum(1 for r in records if "E2E" in r["Loại"])
    unm = len(by.get("UNMAPPED", []))
    L.append(f"**Tổng FO: {total} TC** — E2E: {e2e} · Edge case: {total - e2e - unm} · Chưa phân loại: {unm}\n")

    L.append("## 1. Vai trò & trạng thái\n")
    L.append("| Vai trò | Dùng ở bước |\n|---|---|")
    L.append("| Maker | Tạo mới, Lưu nháp, Gửi duyệt, Chỉnh sửa, Hủy |")
    L.append("| Kế toán phó | Duyệt cấp 1 (luồng 2 cấp), 2FA |")
    L.append("| Kế toán trưởng | Duyệt (cấp cuối), 2FA, Trả lại, Từ chối |")
    L.append("| Chủ tài khoản | Ký số cá nhân, Từ chối |")
    L.append("| Người ký pháp nhân | Ký số pháp nhân (không có Từ chối) |\n")
    L.append("```")
    L.append("Khởi tạo phát hành ─(Lưu nháp)→ Bản nháp ─(Gửi duyệt)→ Chờ duyệt")
    L.append("Chờ duyệt ─(KT phó duyệt, luồng 2 cấp)→ Chờ duyệt cấp tiếp theo ─(KT trưởng duyệt)→ Chờ duyệt (1 cấp) / Chờ duyệt cấp tiếp theo (2 cấp)")
    L.append("         ─(Chủ TK ký số)→ [Chờ ký số pháp nhân ─(ký PN)→] Đã gửi ngân hàng  ⇒ chuyển sang BO")
    L.append("Nhánh lỗi: Trả lại → Yêu cầu chỉnh sửa · KT trưởng từ chối → Đã hủy · Chủ TK từ chối → Từ chối · Maker hủy → Đã hủy")
    L.append("```\n")

    L.append("## 2. Luồng happy (E2E)\n")
    L.append("Mỗi luồng có 4 biến thể: số cấp duyệt (1 | 2) × ký số pháp nhân (không | có).\n")
    L.append("| Flow | Mô tả | 1 cấp, không PN | 1 cấp, có PN | 2 cấp, không PN | 2 cấp, có PN |\n|---|---|---|---|---|---|")
    for fid, desc, variants in FLOWS:
        inv = {v: k for k, v in variants.items()}
        L.append(f"| **{fid}** | {desc} | TC{inv['1 cấp, không ký số PN']} | TC{inv['1 cấp, có ký số PN']} | "
                 f"TC{inv['2 cấp, không ký số PN']} | TC{inv['2 cấp, có ký số PN']} |")
    L.append("\n**Luồng lỗi (E2E):**\n")
    for no, d in NEGATIVE_E2E.items():
        L.append(f"- TC{no} – {d}")
    L.append("")

    L.append("## 3. Edge case theo từng bước của luồng\n")
    L.append("Thứ tự bước theo đúng trình tự người dùng đi trong luồng happy. "
             "Cột *Ưu tiên* = số TC Critical/Highest · High.\n")
    sections = [("Bước chính của luồng happy", "S"), ("Nhánh lỗi", "N"), ("Tiện ích", "U"),
                ("Sau phát hành (phụ thuộc BO)", "P")]
    for title, prefix in sections:
        L.append(f"### {title}\n")
        L.append("| Step | Bước | Áp dụng | Số TC | Ưu tiên | TC No |\n|---|---|---|---|---|---|")
        for sid, name, applies, _ in STEPS:
            if not sid.startswith(prefix) or sid not in by:
                continue
            recs = by[sid]
            hi, h = prio(recs)
            L.append(f"| {sid} | {name} | {applies} | {len(recs)} | {hi} · {h} | {ids_compact(tc_no(r['TC No']) for r in recs)} |")
        L.append("")
    if unm:
        L.append("### ⚠ Chưa phân loại\n")
        L.append(", ".join(f"TC{r['TC No']}" for r in by["UNMAPPED"]) + "\n")

    L.append("## 4. Thứ tự triển khai đề xuất\n")
    L.append("1. **HF-01, biến thể 1 cấp không ký số PN (TC1)** – xương sống, dựng sẵn keyword cho login theo vai trò, "
             "tạo LC tối thiểu, gửi duyệt, phê duyệt 2FA, ký số.")
    L.append("2. Mở rộng 3 biến thể còn lại của HF-01 (TC2–4) → chỉ thêm bước duyệt cấp 2 và ký số PN.")
    L.append("3. HF-02 (S16 Lưu nháp + S17 Chỉnh sửa) và HF-03 (Trả lại).")
    L.append("4. Luồng lỗi E2E (TC9–13).")
    L.append("5. Edge case: ưu tiên S02, S16, S19–S21 (hành vi luồng), sau đó đến các section field S03–S15 "
             "(data-driven, 1 test case Katalon + Data File cho mỗi section).")
    L.append("6. P01–P02 làm sau khi có luồng BO phát hành.\n")

    L.append("## 5. Câu hỏi mở / cần xác nhận\n")
    L.append("- [ ] Ký số cá nhân/pháp nhân trên SIT dùng gì (USB token, HSM, mock)? Có automate được không?")
    L.append("- [ ] 2FA trên SIT: OTP cố định / bypass / phải lấy từ SMS-email?")
    L.append("- [ ] Bộ user test cho từng vai trò (Maker, KT phó, KT trưởng, Chủ TK, Người ký PN) + cấu hình 1 cấp / 2 cấp / có-không ký số PN.")
    L.append("- [ ] Dữ liệu tích hợp (CIF, hợp đồng hạn mức GLIMS, Core) cho S02 có ổn định trên SIT không?")
    L.append("- [ ] Ý nghĩa Note = \"Phục hồi\" trong file gốc (còn trong phạm vi test không?).")
    L.append("- [ ] Ranh giới section S03–S15 được suy ra theo thứ tự TC + tên trường, chưa đối chiếu SRS – cần review.\n")
    FLOWS_MD.write_text("\n".join(L), encoding="utf-8")


def main():
    src = Path(sys.argv[1])
    records = []
    for r in load_fo_rows(src):
        sid, name, kind, applies = classify(r)
        records.append({
            "TC No": tc_no(r[0]) or r[0], "Type": r[1] or "FO-E2E", "Flow / Step ID": sid, "Tên bước": name,
            "Áp dụng cho luồng": applies, "Loại": kind, "Màn hình (gốc)": r[4], "Function": r[2],
            "Điều kiện 1": r[5], "Điều kiện 2": r[6], "Summary": r[8], "Data": r[9], "Action": r[10],
            "Expected Result": r[11], "Priority": r[12], "Gợi ý automation": hint(r), "Note (gốc)": r[17],
        })
    write_xlsx(records)
    write_md(records, src.name)
    unm = [r["TC No"] for r in records if r["Flow / Step ID"] == "UNMAPPED"]
    print(f"FO rows: {len(records)} · unmapped: {len(unm)} {unm[:20]}")


if __name__ == "__main__":
    main()
